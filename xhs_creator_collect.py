#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
小红书创作者中心 · 自有账号笔记数据采集器
==========================================

设计原则
--------
**不碰签名算法。** 做法是：用持久化浏览器 profile 保住登录态，打开创作者中心
「数据看板 → 笔记数据」，让页面自己去发那些带签名的 XHR，我们只在旁边监听
响应、把 JSON 抄下来。

为什么不用开放平台 API：小红书开放平台没有给专业号开放笔记数据接口。
为什么不用第三方数据接口（yddm 等）：字段残缺（无曝光/观看），且按次计费。

四个子命令
----------
  login    扫码登录，把登录态存进 profile（只需做一次，失效了再做）
  check    检查登录态是否还有效（放 cron 里当健康检查）
  probe    探针：打开数据看板，把这段时间内所有 XHR 的 URL 和 JSON 结构打印出来
           —— **第一次接入必须先跑这个**，确认接口和字段名
  collect  正式采集：翻页抓全部笔记 → 本地按日期过滤 → 导出 Excel + CSV

典型用法
--------
  python xhs_creator_collect.py login   --qr-out /tmp/xhs_qr.png
  python xhs_creator_collect.py probe   --dashboard-url "<从地址栏复制>"
  python xhs_creator_collect.py collect --start 2026-06-01 --end 2026-10-31 \\
                                        --out notes.xlsx

合规边界（务必读一遍）
----------------------
这是"用本账号登录态读本账号后台"，只用于自家经营分析。不要对外提供数据服务，
不要扩到他人账号，不要并发、不要高频。详见 README.md。
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import csv
import json
import os
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

# Windows 控制台默认 GBK，中文会乱码；Linux 上无害
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# ---------------------------------------------------------------- 常量

HOME_URL = "https://creator.xiaohongshu.com/"
# 实测（2026-10-29）：逐篇笔记数据页就是这条路由，不用再去点侧边栏
# 数据看板 → 内容分析 那套 d-sub-menu（它是 SPA，没有 <a href>，点文字不跳转）。
DEFAULT_DASHBOARD_URL = "https://creator.xiaohongshu.com/statistics/data-analysis"
# 一个尽量像真人的 UA。Playwright 默认 UA 里带 HeadlessChrome，必须换掉。
REAL_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)

# 用来"认领"一条笔记对象的键
ID_KEYS = {"noteId", "note_id", "noteID", "id", "note_id_str"}
# 出现这些键，说明这条列表大概率是笔记数据
HINT_KEYS = {
    "viewCount", "views", "readCount", "impression", "impressions",
    "likedCount", "likes", "likeCount", "collectedCount", "collectCount",
    "commentCount", "commentsCount", "shareCount", "share_count",
    "view_count", "liked_count", "collected_count", "comments_count",
}

# 字段别名表：页面里的名字 → 我们输出的列名
ALIASES: dict[str, list[str]] = {
    "note_id": ["noteId", "note_id", "noteID", "noteIdStr", "id", "note_id_str"],
    "title": ["title", "displayTitle", "display_title", "noteTitle", "name"],
    "publish_time": [
        "time", "createTime", "create_time", "publishTime", "publish_time",
        "firstPublishTime", "noteFirstPublishTime", "first_publish_time",
        "publishDate", "postTime", "post_time",
    ],
    "type": ["type", "noteType", "note_type", "mediaType"],
    "impressions": [
        "impression", "impressions", "impressionCount", "impression_count",
        "exposure", "exposureCount", "exposeCount",
    ],
    "views": [
        "viewCount", "views", "readCount", "view_count", "read_count",
        "watchCount", "watch_count", "playCount", "play_count",
    ],
    "liked_count": [
        "likedCount", "liked_count", "likes", "likeCount", "like_count",
        "likedNum", "liked",
    ],
    "collected_count": [
        "collectedCount", "collected_count", "collectCount", "collect_count",
        "favCount", "fav_count", "collected", "favouredCount",
    ],
    "comments_count": [
        "commentCount", "comment_count", "commentsCount", "comments_count",
        "commentNum", "comment",
    ],
    "share_count": ["shareCount", "share_count", "shares", "sharedCount", "shareNum"],
    "follows": [
        "followCount", "follow_count", "follows", "addFansCount",
        "fansIncrease", "newFans", "increaseFansCount", "growFansCount",
    ],
    "ctr": ["cover_click_rate", "ctr", "click_rate", "coverClickRate"],
    "avg_watch_sec": ["avg_watch_time", "avg_watch", "per_capita_watch_time"],
    "danmaku": ["danmaku_count", "danmaku"],
    "audit_status": ["audit_status"],
}

# 官方「导出数据」表头 → 我们的列名（实测 2026-10-29）
OFFICIAL_MAP = {
    "笔记标题": "title",
    "首次发布时间": "publish_time",
    "体裁": "type",
    "曝光": "impressions",
    "观看量": "views",
    "封面点击率": "ctr",
    "点赞": "liked_count",
    "评论": "comments_count",
    "收藏": "collected_count",
    "涨粉": "follows",
    "分享": "share_count",
    "人均观看时长": "avg_watch_sec",
    "弹幕": "danmaku",
}

# 输出列（和上一版 xls 的列名尽量对齐，下游脚本不用改）
COLUMNS = [
    "user_id", "note_id", "title", "publish_time", "type",
    "liked_count", "collected_count", "comments_count", "share_count",
    "impressions", "views", "ctr", "follows", "avg_watch_sec", "danmaku",
    "fetch_time",
]
NUMERIC_COLUMNS = {
    "liked_count", "collected_count", "comments_count", "share_count",
    "impressions", "views", "ctr", "follows", "avg_watch_sec", "danmaku",
    "user_id",
}


# ---------------------------------------------------------------- 小工具

def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", file=sys.stderr, flush=True)


def human_delay(base: float) -> None:
    """比真人略保守的抖动间隔。"""
    time.sleep(max(0.2, base + random.uniform(-0.6, 1.2)))


def to_number(v):
    """'1.2万' -> 12000；'--' / '' -> None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip()
    if not s or s in {"-", "--", "—", "null", "None", "N/A"}:
        return None
    mult = 1
    if s.endswith(("万", "w", "W")):
        mult, s = 10000, s[:-1]
    elif s.endswith("亿"):
        mult, s = 100000000, s[:-1]
    s = s.replace(",", "").replace("+", "").replace("%", "").strip()
    try:
        x = float(s) * mult
    except ValueError:
        return None
    return int(x) if float(x).is_integer() else x


def _norm_type(v) -> str:
    """接口里 type 是数字（1=图文 / 2=视频），导出表里是中文，统一成中文。"""
    s = str(v).strip()
    return {"1": "图文", "2": "视频", "3": "视频",
            "normal": "图文", "video": "视频"}.get(s.lower(), s) or s


def parse_time(v):
    """各种时间写法 -> 'YYYY-MM-DD HH:MM:SS'；解析不出来就原样返回。"""
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit()):
        n = int(float(v))
        if n > 10**12:      # 毫秒
            n //= 1000
        if 10**9 < n < 10**11:
            return datetime.fromtimestamp(n).strftime("%Y-%m-%d %H:%M:%S")
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                "%Y/%m/%d %H:%M:%S", "%Y/%m/%d", "%m-%d %H:%M", "%m-%d"):
        try:
            dt = datetime.strptime(s, fmt)
            if fmt.startswith("%m"):
                dt = dt.replace(year=datetime.now().year)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
        return datetime(y, mo, d).strftime("%Y-%m-%d %H:%M:%S")
    return s


def strip_query(url: str) -> str:
    return url.split("?", 1)[0]


# ---------------------------------------------------------------- JSON 结构识别

def iter_note_lists(obj, path="$", depth=0, out=None):
    """递归找出"看起来是笔记列表"的数组（元素是含 id + 互动字段的 dict）。"""
    if out is None:
        out = []
    if depth > 5:
        return out
    if isinstance(obj, dict):
        for k, v in obj.items():
            iter_note_lists(v, f"{path}.{k}", depth + 1, out)
    elif isinstance(obj, list) and obj and isinstance(obj[0], dict):
        keys = set(obj[0].keys())
        if (keys & ID_KEYS) and (keys & HINT_KEYS):
            out.append({"path": path, "items": obj, "keys": sorted(keys)})
    return out


def score_candidate(cand: dict) -> int:
    """数量 * 命中字段数，用来在多个候选里挑最像"笔记列表"的那个。"""
    keys = set(cand["keys"])
    hits = len(keys & HINT_KEYS) + 2 * len(keys & ID_KEYS)
    return hits * max(1, len(cand["items"]))


def extract(payload: dict) -> dict:
    """按别名表把一条原始笔记 dict 映射成输出行。"""
    lower = {str(k).lower(): v for k, v in payload.items()}
    row: dict[str, object] = {}
    for col, names in ALIASES.items():
        val = None
        for n in names:
            if n in payload and payload[n] not in (None, ""):
                val = payload[n]
                break
            if n.lower() in lower and lower[n.lower()] not in (None, ""):
                val = lower[n.lower()]
                break
        if val is None:
            continue
        if col == "publish_time":
            row[col] = parse_time(val)
        elif col in NUMERIC_COLUMNS:
            row[col] = to_number(val)
        elif col == "type":
            row[col] = _norm_type(val)
        else:
            row[col] = str(val).strip()
    # 补齐所有列（没有的填 None），保证下游拿到的列是稳定的
    for col in COLUMNS:
        row.setdefault(col, None)
    return row


# ---------------------------------------------------------------- 浏览器

async def launch(p, args):
    """持久化 profile 启动。一个账号一个目录，绝不复用。"""
    profile = Path(args.profile).expanduser()
    profile.mkdir(parents=True, exist_ok=True)
    argv = [
        "--disable-blink-features=AutomationControlled",
        "--disable-dev-shm-usage",
        "--no-first-run",
        "--no-default-browser-check",
        "--disable-features=IsolateOrigins,site-per-process",
    ]
    if args.no_sandbox:
        argv.append("--no-sandbox")
    ctx = await p.chromium.launch_persistent_context(
        user_data_dir=str(profile),
        headless=not args.headful,
        args=argv,
        viewport={"width": 1440, "height": 900},
        locale="zh-CN",
        timezone_id="Asia/Shanghai",
        user_agent=REAL_UA,
        accept_downloads=True,
    )
    await ctx.add_init_script(
        "Object.defineProperty(navigator,'webdriver',{get:()=>undefined});"
        "Object.defineProperty(navigator,'languages',{get:()=>['zh-CN','zh']});"
        "Object.defineProperty(navigator,'plugins',{get:()=>[1,2,3,4,5]});"
    )
    return ctx


LOGIN_TEXTS = ("扫码登录", "立即登录", "登录小红书", "手机号登录", "验证码登录")
# 实际 cookie 名带域名后缀，例如 access-token-creator.xiaohongshu.com，
# 所以用子串匹配而不是精确集合。
# 实测坑：未登录的游客也会带 customer-sso-sid / galaxy.creator.beaker.session.id，
# 所以 "session"/"sso" 这种宽泛关键词会造成误判。只用真正的登录态 cookie。
SESSION_COOKIE_MARKERS = ("access-token", "galaxy_creator_session", "x-user-id-creator")


async def has_session_cookie(ctx) -> bool:
    try:
        names = [str(c.get("name", "")).lower() for c in await ctx.cookies()]
        return any(m in n for n in names for m in SESSION_COOKIE_MARKERS)
    except Exception:
        return False


async def is_logged_in(page, ctx) -> bool:
    """只有拿到**正向证据**才算登录。

    踩过的坑：切到扫码模式后 SPA 会把 URL 里的 login 去掉，如果按
    "URL 不含 login 就算登录"来判，会在二维码还没扫的瞬间误判成功。
    所以这里要求：登录态 cookie，或者页面上出现了「小红书账号: xxx」。
    """
    if "login" in page.url.lower():
        return False
    if await has_session_cookie(ctx):
        return True
    if await read_red_id(page):
        return True
    return False


QR_SELECTORS = [
    "img[src*='qrcode']", "img[src*='qr']", ".qrcode img", ".qrcode-img",
    ".login-qrcode img", "canvas.qrcode", "canvas", ".login-container img",
    ".login-container canvas",
]

# 登录页的类名是哈希过的（css-1lhmg90 这种），靠选择器认不住，
# 所以这里按"几何 + 是否 data URI"来找二维码：方形、110~340px、内嵌图片。
_FIND_QR_JS = """() => {
  const els = [...document.querySelectorAll('img,canvas')];
  let best = null, bestScore = -1;
  for (const el of els) {
    const r = el.getBoundingClientRect();
    const w = Math.round(r.width), h = Math.round(r.height);
    if (w < 110 || h < 110 || w > 340 || h > 340) continue;
    const ratio = w / h;
    if (ratio < 0.85 || ratio > 1.18) continue;
    let uri = '', score = w;
    if (el.tagName === 'IMG') {
      uri = el.currentSrc || el.src || '';
      if (/^data:image/.test(uri)) score += 300;
    } else {
      try { uri = el.toDataURL('image/png'); score += 300; } catch (e) { uri = ''; }
    }
    if (!uri) continue;
    if (score > bestScore) {
      bestScore = score;
      best = {uri: uri, x: Math.round(r.x), y: Math.round(r.y), w: w, h: h};
    }
  }
  return best;
}"""


def _save_data_uri(uri: str, out_path: Path) -> bool:
    m = re.match(r"data:image/[a-zA-Z+]+;base64,(.*)", uri, re.S)
    if not m:
        return False
    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_bytes(base64.b64decode(m.group(1)))
        return True
    except Exception:
        return False


async def _click_qr_toggle(page) -> bool:
    """把登录卡从「短信登录」切到「扫码登录」。切不过去就返回 False。"""
    for text in ("扫码登录", "二维码登录"):
        try:
            loc = page.get_by_text(text, exact=True).first
            if await loc.count() > 0 and await loc.is_visible():
                await loc.click(timeout=2500)
                await page.wait_for_timeout(1500)
                return True
        except Exception:
            continue
    # 兜底：登录卡右上角那个小方形图标
    try:
        box = await page.evaluate("""() => {
          const cands = [];
          for (const el of document.querySelectorAll('img,canvas,svg,div')) {
            const r = el.getBoundingClientRect();
            if (r.width < 28 || r.width > 100 || r.x < 900) continue;
            if (Math.abs(r.width - r.height) > 12) continue;
            cands.push({x: r.x + r.width / 2, y: r.y + r.height / 2});
          }
          if (!cands.length) return null;
          cands.sort((a, b) => a.y - b.y);
          return cands[0];
        }""")
        if box:
            await page.mouse.click(box["x"], box["y"])
            await page.wait_for_timeout(1500)
            return True
    except Exception:
        pass
    return False


async def _wait_qr(page, timeout_ms: int = 15000):
    deadline = time.time() + timeout_ms / 1000.0
    while time.time() < deadline:
        try:
            qr = await page.evaluate(_FIND_QR_JS)
        except Exception:
            qr = None
        if qr:
            return qr
        await page.wait_for_timeout(1000)
    return None


async def capture_qr(page, out_path: Path) -> Path:
    """把登录二维码原图存成 PNG。

    登录页类名是哈希的（css-1lhmg90），选择器不可靠，所以直接用页面里的
    data URI 解码，拿到的是原图（通常 320px），比截 160px 的缩略图清楚得多。

    实测坑：这个页面默认停在「短信登录」，二维码面板要么延迟挂载、要么藏在
    侧边，所以先判断模式，再点切换，然后轮询等它出现。
    """
    out_path = Path(out_path)

    # 如果停在短信登录，先切到扫码
    try:
        sms = page.get_by_text("短信登录", exact=False).first
        if await sms.count() > 0 and await sms.is_visible():
            log("   登录页停在短信模式，切到扫码…")
            await _click_qr_toggle(page)
    except Exception:
        pass

    qr = await _wait_qr(page, 18000)
    if qr is None:
        log("   还没出码，再点一次切换…")
        await _click_qr_toggle(page)
        qr = await _wait_qr(page, 8000)

    if qr and _save_data_uri(qr["uri"], out_path):
        log(f"   二维码原图 {qr['w']}x{qr['h']} @ ({qr['x']},{qr['y']})")
        return out_path

    # 兜底：整页截图（至少能看见二维码）
    out_path.parent.mkdir(parents=True, exist_ok=True)
    await page.screenshot(path=str(out_path), full_page=False)
    log("   ⚠ 没取到二维码原图，退回整页截图")
    return out_path


async def goto_dashboard(page, args) -> str:
    """打开逐篇笔记数据页。

    实测：数据看板是 d-sub-menu（SPA，没有 <a href>），点标题不会跳转；
    但 /statistics/data-analysis 这条路由可以直接 goto，最稳。
    """
    url = args.dashboard_url or DEFAULT_DASHBOARD_URL
    await page.goto(url, wait_until="domcontentloaded")
    await page.wait_for_timeout(3500)
    return page.url


# ---------------------------------------------------------------- 响应捕获

class Capture:
    def __init__(self, raw_dir: Path | None, pin: str | None, max_bytes: int = 8 << 20):
        self.raw_dir = raw_dir
        self.pin = pin
        self.max_bytes = max_bytes
        self.bodies: list[dict] = []      # {"url","path","items":[...],"score":int}
        self.all_urls: list[str] = []     # probe 用：所有见过的 JSON 接口
        self._seen: set[str] = set()

    def interesting(self, url: str) -> bool:
        if self.pin and self.pin not in url:
            return False
        return "/api/" in url or url.endswith(".json")

    async def on_response(self, resp):
        try:
            url = resp.url
            if not self.interesting(url):
                return
            ctype = (resp.headers or {}).get("content-type", "")
            if "json" not in ctype and not url.endswith(".json"):
                return
            if url in self._seen:
                return
            self._seen.add(url)
            self.all_urls.append(url)
            body = await resp.json()
        except Exception:
            return

        cands = iter_note_lists(body)
        if not cands:
            return
        best = max(cands, key=score_candidate)
        self.bodies.append({
            "url": url,
            "path": strip_query(url),
            "items": best["items"],
            "keys": best["keys"],
            "score": score_candidate(best),
        })
        log(f"  ↳ 命中 {strip_query(url)}  {len(best['items'])} 条  keys={best['keys'][:8]}")

        if self.raw_dir:
            self.raw_dir.mkdir(parents=True, exist_ok=True)
            name = re.sub(r"[^A-Za-z0-9_.-]", "_", strip_query(url).rsplit("/api/", 1)[-1])[:80]
            fp = self.raw_dir / f"{name}__{len(self.bodies):03d}.json"
            try:
                fp.write_text(json.dumps(body, ensure_ascii=False), encoding="utf-8")
            except Exception:
                pass

    def merged(self) -> tuple[str, list[dict], list[str]]:
        """把所有响应按接口路径分组、合并、去重，挑出最像"自有笔记数据"的那组。

        踩过的坑：创作学院的推荐卡接口（create_guidance）也会返回 note_id +
        view_count，条数还比真实笔记多。不能按条数挑，要按"含发布时间"挑 ——
        创作学院的卡片没有 post_time，自己的笔记才有。
        """
        groups: dict[str, list[dict]] = {}
        for b in self.bodies:
            groups.setdefault(b["path"], []).extend(b["items"])
        if not groups:
            return "", [], []

        best_path, best_score = "", -1
        for path, items in groups.items():
            with_time = 0
            for it in items[:60]:
                row = extract(it)
                if row.get("note_id") and row.get("publish_time"):
                    with_time += 1
            score = with_time * 10 + len(items)
            if score > best_score:
                best_score, best_path = score, path

        items = groups[best_path]
        seen: set[str] = set()
        uniq: list[dict] = []
        for it in items:
            row = extract(it)
            nid = str(row.get("note_id") or "")
            if not nid:
                continue
            if nid in seen:
                continue
            seen.add(nid)
            uniq.append(row)
        return best_path, uniq, sorted({k for it in items[:3] for k in it})


async def goto_next_page(page) -> bool:
    """点「下一页」。找不到或已禁用就返回 False。"""
    selectors = [
        "button:has-text('下一页')",
        "[role='button']:has-text('下一页')",
        "li:has-text('下一页')",
        "a:has-text('下一页')",
        ".ant-pagination-next",
        ".pagination-next",
        "text=下一页",
    ]
    for sel in selectors:
        try:
            loc = page.locator(sel).first
            if await loc.count() == 0 or not await loc.is_visible():
                continue
            attrs = " ".join([
                (await loc.get_attribute("class") or ""),
                (await loc.get_attribute("aria-disabled") or ""),
                (await loc.get_attribute("disabled") or ""),
            ]).lower()
            if "disabled" in attrs:
                return False
            await loc.scroll_into_view_if_needed(timeout=3000)
            await loc.click(timeout=6000)
            return True
        except Exception:
            continue
    return False


_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


async def _date_boxes(page):
    try:
        boxes = await page.evaluate("""() => [...document.querySelectorAll('input')]
          .map(i => { const r = i.getBoundingClientRect();
                      return {v: (i.value || '').trim(), x: r.x, y: r.y,
                              w: r.width, h: r.height}; })
          .filter(o => o.w > 40 && o.h > 8)""")
    except Exception:
        return []
    return [b for b in boxes if _DATE_RE.match(b["v"])][:2]


async def try_set_range(page, start: str, end: str) -> bool:
    """填「笔记首发时间」的开始/结束。

    实测这个页面默认只查近 30 天，不填日期就会漏掉更早的笔记。
    """
    for idx, val in enumerate((start, end)):
        boxes = await _date_boxes(page)
        if len(boxes) < 2:
            return False
        box = boxes[idx]
        try:
            await page.mouse.click(box["x"] + box["w"] / 2, box["y"] + box["h"] / 2)
            await page.wait_for_timeout(500)
            await page.keyboard.press("Control+A")
            await page.keyboard.type(val, delay=60)
            await page.keyboard.press("Enter")
            await page.wait_for_timeout(900)
        except Exception:
            return False
    return True


async def read_red_id(page) -> str:
    """从页面文本里读「小红书账号: xxxx」。"""
    try:
        txt = await page.evaluate("() => document.body.innerText || ''")
    except Exception:
        return ""
    m = re.search(r"小红书账号[:：]\s*([0-9A-Za-z_]+)", txt)
    return m.group(1) if m else ""


async def try_export(page, out_dir: Path):
    """点官方「导出数据」，把下载的文件存下来。"""
    try:
        loc = page.get_by_text("导出数据", exact=False).first
        if await loc.count() == 0 or not await loc.is_visible():
            log("   页面上没找到「导出数据」按钮")
            return None
        out_dir.mkdir(parents=True, exist_ok=True)
        async with page.expect_download(timeout=25000) as dl_info:
            await loc.click(timeout=5000)
        dl = await dl_info.value
        target = out_dir / dl.suggested_filename
        await dl.save_as(str(target))
        return target
    except Exception as e:
        log(f"   官方导出没跑通（{type(e).__name__}），回退到接口抓取")
        return None


_CN_DT_RE = re.compile(r"(\d{4})年(\d{1,2})月(\d{1,2})日(?:(\d{1,2})时(\d{1,2})分(\d{1,2})秒)?")


def _parse_cn_time(v):
    """'2026年10月29日14时42分21秒' -> 'YYYY-MM-DD HH:MM:SS'。"""
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    m = _CN_DT_RE.search(str(v))
    if not m:
        return parse_time(v)
    y, mo, d, hh, mi, ss = (int(x) if x else 0 for x in m.groups())
    return datetime(y, mo, d, hh, mi, ss).strftime("%Y-%m-%d %H:%M:%S")


def parse_official_xlsx(path: Path) -> list[dict]:
    """解析官方「导出数据」下载的 xlsx。

    实测格式：第 1 行是提示「最多导出排序后前1000条笔记」，第 2 行才是表头，
    第 3 行起是数据，且 **不含笔记 ID**。
    """
    try:
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        log(f"   ⚠ 读不了官方导出表: {type(e).__name__}")
        return []

    hi = None
    for i, r in enumerate(rows[:6]):
        if r and any(isinstance(c, str) and "笔记标题" in c for c in r):
            hi = i
            break
    if hi is None:
        log("   ⚠ 官方导出表里找不到表头「笔记标题」，放弃解析")
        return []

    headers = [str(c).strip() if c is not None else "" for c in rows[hi]]
    out: list[dict] = []
    for r in rows[hi + 1:]:
        if not any(c not in (None, "") for c in r):
            continue
        raw = dict(zip(headers, r))
        row = {col: None for col in COLUMNS}
        for cn, val in raw.items():
            col = OFFICIAL_MAP.get(cn)
            if not col or val in (None, ""):
                continue
            if col == "publish_time":
                row[col] = _parse_cn_time(val)
            elif col in NUMERIC_COLUMNS:
                row[col] = to_number(val)
            else:
                row[col] = str(val).strip()
        if row.get("publish_time") or row.get("liked_count") is not None:
            out.append(row)
    return out


# ---------------------------------------------------------------- 输出

def write_outputs(rows: list[dict], out_path: Path, user_id: str) -> tuple[Path, Path]:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for r in rows:
        r["user_id"] = user_id
        r["fetch_time"] = now
    out_path.parent.mkdir(parents=True, exist_ok=True)
    csv_path = out_path.with_suffix(".csv")

    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "笔记数据"
        ws.append(COLUMNS)
        head_fill = PatternFill("solid", fgColor="F2F2F2")
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append([r.get(c, "") for c in COLUMNS])
        for idx, col in enumerate(COLUMNS, start=1):
            letter = get_column_letter(idx)
            width = 26 if col == "title" else (20 if col == "note_id" else 14)
            ws.column_dimensions[letter].width = width
            if col in NUMERIC_COLUMNS:
                for cell in ws[letter][1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
        wb.save(out_path)
    except ImportError:
        log("⚠ 未装 openpyxl，只输出 CSV")
        out_path = csv_path

    return out_path, csv_path


# ---------------------------------------------------------------- 子命令

async def cmd_login(args) -> int:
    async with async_playwright_ctx() as pw:
        ctx = await launch(pw, args)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        await page.goto(HOME_URL, wait_until="domcontentloaded")
        await page.wait_for_timeout(3000)
        if await is_logged_in(page, ctx):
            log(f"✅ 已是登录状态（账号 {await read_red_id(page) or '未知'}），profile: {args.profile}")
            await ctx.close()
            return 0

        qr = await capture_qr(page, Path(args.qr_out).expanduser())
        log(f"📷 二维码已保存: {qr}")
        log("   把这张图给人扫（手机小红书 → 我的 → 扫一扫），或者 scp 下来看。")
        deadline = time.time() + args.timeout
        while time.time() < deadline:
            await page.wait_for_timeout(2000)
            if await is_logged_in(page, ctx):
                log(f"✅ 登录成功（账号 {await read_red_id(page) or '未知'}），登录态已写入 profile")
                await page.wait_for_timeout(3000)
                await ctx.close()
                return 0
        log(f"❌ 等待 {args.timeout}s 未完成登录（可能二维码已过期）")
        await ctx.close()
        return 1


async def cmd_check(args) -> int:
    async with async_playwright_ctx() as pw:
        ctx = await launch(pw, args)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        await page.goto(HOME_URL, wait_until="domcontentloaded")
        await page.wait_for_timeout(3000)
        ok = await is_logged_in(page, ctx)
        red_id = await read_red_id(page) if ok else ""
        print(json.dumps({
            "logged_in": ok,
            "red_id": red_id,
            "url": page.url,
            "has_session_cookie": await has_session_cookie(ctx),
            "profile": str(Path(args.profile).expanduser()),
        }, ensure_ascii=False, indent=2))
        await ctx.close()
        return 0 if ok else 3


async def cmd_probe(args) -> int:
    """第一次接入必跑：告诉你接口在哪、字段叫什么。"""
    capture = Capture(raw_dir=Path(args.raw_dir).expanduser() if args.keep_raw else None,
                      pin=args.pin_url)
    async with async_playwright_ctx() as pw:
        ctx = await launch(pw, args)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        page.on("response", capture.on_response)
        try:
            url = await goto_dashboard(page, args)
        except SystemExit:
            await ctx.close()
            raise
        log(f"📍 当前页面: {url}")
        log(f"⏳ 监听 {args.seconds}s ……")
        await page.wait_for_timeout(int(args.seconds * 1000))

        report = {
            "page_url": page.url,
            "logged_in": await is_logged_in(page, ctx),
            "json_endpoints_seen": sorted(set(strip_query(u) for u in capture.all_urls)),
            "note_list_candidates": [
                {
                    "url": b["url"],
                    "path": b["path"],
                    "count": len(b["items"]),
                    "score": b["score"],
                    "sample_keys": b["keys"],
                    "sample": {k: b["items"][0].get(k) for k in b["keys"][:12]},
                }
                for b in capture.bodies
            ],
        }
        out = Path(args.out).expanduser()
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

        print(json.dumps(report, ensure_ascii=False, indent=2))
        shot = out.with_name("probe_page.png")
        try:
            await page.screenshot(path=str(shot))
            log(f"🖼  页面截图: {shot}")
        except Exception:
            pass
        log(f"📄 探针报告: {out}")
        if not capture.bodies:
            log("⚠ 没抓到笔记列表。可能原因：① 登录态失效 ② 页面没有真正打开数据看板 "
                "③ 数据是服务端渲染（那就得改抓 DOM）—— 把报告发我看看")
        await ctx.close()
        return 0


async def cmd_collect(args) -> int:
    if not args.skip_check:
        rc = await cmd_check(args)
        if rc != 0:
            log("❌ 登录态失效，先跑: python xhs_creator_collect.py login")
            return 3

    raw_dir = Path(args.raw_dir).expanduser() if args.keep_raw else None
    capture = Capture(raw_dir=raw_dir, pin=args.pin_url)

    async with async_playwright_ctx() as pw:
        ctx = await launch(pw, args)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        page.on("response", capture.on_response)

        # 先回首页读小红书号（数据页上不显示），用作 user_id 列
        await page.goto(HOME_URL, wait_until="domcontentloaded")
        await page.wait_for_timeout(2500)
        red_id = args.user_id or await read_red_id(page)
        log(f"   账号: {red_id or '（未识别，user_id 列留空，可用 --user-id 指定）'}")

        url = await goto_dashboard(page, args)
        log(f"📍 数据看板: {url}")
        await page.wait_for_timeout(4000)

        if args.start and args.end:
            ok = await try_set_range(page, args.start, args.end)
            log("   时间筛选: " + (f"已填入 {args.start} ~ {args.end}" if ok
                                   else "没找到日期输入框，只做本地过滤"))
            await page.wait_for_timeout(2500)

        # 先翻页把接口数据全抓下来（用途：补 note_id —— 官方导出表里没有笔记 ID）
        pages = 1
        while pages < args.max_pages:
            before = len(capture.bodies)
            if not await goto_next_page(page):
                log(f"🚩 第 {pages} 页无「下一页」，停止翻页")
                break
            for _ in range(20):
                await asyncio.sleep(0.4)
                if len(capture.bodies) > before:
                    break
            pages += 1
            log(f"   翻到第 {pages} 页")
            human_delay(args.delay)

        api_path, api_rows, keys = capture.merged()
        log(f"📥 接口 {api_path or '(未命中)'}  抓到 {len(api_rows)} 条（去重后）")

        # 优先走官方导出：字段最全（曝光 / 封面点击率 / 人均观看时长 / 弹幕）
        rows: list[dict] = []
        if args.export:
            exported = await try_export(page, Path(args.out).expanduser().parent)
            if exported:
                log(f"✅ 官方导出成功: {exported}")
                rows = parse_official_xlsx(exported)
                by_time = {r["publish_time"]: r for r in api_rows
                           if r.get("publish_time") and r.get("note_id")}
                hit = 0
                for r in rows:
                    m = by_time.get(r.get("publish_time"))
                    if m:
                        r["note_id"] = m["note_id"]
                        hit += 1
                log(f"   官方表 {len(rows)} 行，接口补到 note_id {hit} 条")

        if not rows:
            rows = api_rows
            if not rows:
                if not await has_session_cookie(ctx):
                    log("❌ 没有 session cookie，登录态多半已经失效。先跑 login。")
                else:
                    log("❌ 没采到任何笔记。先跑 probe 看接口结构。")
                await ctx.close()
                return 4
            log("⚠ 回退到接口数据：缺 曝光/封面点击率/涨粉/人均观看时长/弹幕")

        # 本地按发布日期过滤（比在页面上点日期选择器稳得多）
        if args.start or args.end:
            def keep(r):
                d = (r.get("publish_time") or "")[:10]
                if not d:
                    return True
                if args.start and d < args.start:
                    return False
                if args.end and d > args.end:
                    return False
                return True
            before_n = len(rows)
            rows = [r for r in rows if keep(r)]
            log(f"🗓  按 {args.start or '最早'} ~ {args.end or '现在'} 过滤: "
                f"{before_n} → {len(rows)} 条")

        rows.sort(key=lambda r: str(r.get("publish_time") or ""), reverse=True)
        out_path = Path(args.out).expanduser()
        xlsx, csvp = write_outputs(rows, out_path, red_id)
        log(f"✅ 已导出: {xlsx}")
        log(f"✅ 已导出: {csvp}")

        missing = [c for c in ("liked_count", "collected_count", "comments_count") if not any(
            r.get(c) is not None for r in rows)]
        if missing:
            log(f"⚠ 这些列全空，字段名可能变了: {missing}；字段样本: {keys[:20]}")
        if not any(r.get("note_id") for r in rows):
            log("⚠ note_id 一列全空：接口没翻到页，或接口改版。用 发布时间 做联合主键即可")
        await ctx.close()
        return 0


# ---------------------------------------------------------------- 自测

def cmd_selftest(args) -> int:
    """不联网验证"解析 + 导出"这段逻辑。"""
    fake_p1 = {
        "code": 0,
        "data": {
            "list": [
                {"noteId": "aaa1", "title": "第一篇", "time": 1786000000, "type": "normal",
                 "viewCount": "1.2万", "likedCount": "230", "collectedCount": 12,
                 "commentCount": "8", "shareCount": 2, "followCount": 3},
                {"noteId": "aaa2", "title": "第二篇", "time": "2026-07-16 10:00:00",
                 "type": "video", "viewCount": 500, "likedCount": 30,
                 "collectedCount": 4, "commentCount": 1, "shareCount": 0},
            ]
        },
    }
    fake_p2 = {"data": {"notes": [
        {"note_id": "aaa1", "title": "第一篇(重复)", "liked_count": 999},
        {"note_id": "aaa3", "title": "第三篇", "create_time": 1786000000000,
         "liked_count": 7, "collected_count": 1, "comments_count": 0, "share_count": 1},
    ]}}

    rows: list[dict] = []
    seen = set()
    for body in (fake_p1, fake_p2):
        for cand in iter_note_lists(body):
            for it in cand["items"]:
                r = extract(it)
                nid = str(r.get("note_id") or "")
                if nid and nid not in seen:
                    seen.add(nid)
                    rows.append(r)

    assert len(rows) == 3, f"期望 3 条去重后数据，实际 {len(rows)}"
    by_id = {r["note_id"]: r for r in rows}
    assert by_id["aaa1"]["liked_count"] == 230, "重复时应该保留第一次的完整值"
    assert by_id["aaa1"]["views"] == 12000, "'1.2万' 应该转成 12000"
    assert by_id["aaa3"]["publish_time"].startswith("2026-"), by_id["aaa3"]
    assert by_id["aaa3"]["views"] is None, "aaa3 没有观看数"

    tmp = Path(args.out).expanduser().with_suffix(".xlsx")
    xlsx, csvp = write_outputs(rows, tmp, "00000000000")
    print(json.dumps(rows, ensure_ascii=False, indent=2))
    print(f"selftest OK -> {xlsx} / {csvp}")
    return 0


# ---------------------------------------------------------------- 入口

def async_playwright_ctx():
    from playwright.async_api import async_playwright
    return async_playwright()


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="小红书创作者中心 · 自有账号笔记数据采集器",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="先 login → 再 probe → 最后 collect。详见 README.md",
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    def common(sp):
        sp.add_argument("--profile", default="~/.xhs/creator-profile",
                        help="浏览器 profile 目录（一个账号一个）")
        sp.add_argument("--headful", action="store_true",
                        help="有头模式（Linux 上配合 xvfb-run 用，风控更松）")
        sp.add_argument("--no-sandbox", action="store_true",
                        help="root 用户跑 Chromium 时需要")
        sp.add_argument("--dashboard-url", default=os.environ.get("XHS_DASHBOARD_URL"),
                        help="数据看板完整 URL（最稳，从地址栏复制）")
        sp.add_argument("--pin-url", default=None,
                        help="只捕获 URL 含此子串的响应（probe 之后用它锁定接口）")
        sp.add_argument("--keep-raw", action="store_true", help="把原始 JSON 存一份")
        sp.add_argument("--raw-dir", default="raw", help="原始 JSON 目录")

    sp = sub.add_parser("login", help="扫码登录，登录态写入 profile")
    common(sp)
    sp.add_argument("--qr-out", default="out/xhs_qr.png", help="二维码 PNG 输出路径")
    sp.add_argument("--timeout", type=int, default=300, help="等扫码的秒数")
    sp.set_defaults(func=cmd_login)

    sp = sub.add_parser("check", help="检查登录态是否有效")
    common(sp)
    sp.set_defaults(func=cmd_check)

    sp = sub.add_parser("probe", help="探针：列出接口和字段（第一次必跑）")
    common(sp)
    sp.add_argument("--seconds", type=float, default=25, help="监听多少秒")
    sp.add_argument("--out", default="probe_report.json")
    sp.set_defaults(func=cmd_probe)

    sp = sub.add_parser("collect", help="正式采集并导出")
    common(sp)
    sp.add_argument("--start", default=None, help="起始日期 YYYY-MM-DD（本地过滤）")
    sp.add_argument("--end", default=None, help="结束日期 YYYY-MM-DD")
    sp.add_argument("--out", default="xhs_notes.xlsx")
    sp.add_argument("--user-id", default="", help="写进 user_id 列，填你的小红书号或 ID")
    sp.add_argument("--max-pages", type=int, default=50, help="最多翻几页")
    sp.add_argument("--delay", type=float, default=2.0, help="翻页间隔秒数")
    sp.add_argument("--set-range", action="store_true",
                    help="（已默认开启）给了 --start/--end 就会尝试填页面时间筛选")
    sp.add_argument("--export", action="store_true",
                    help="优先点页面上的官方「导出数据」按钮，拿现成的 Excel")
    sp.add_argument("--skip-check", action="store_true", help="跳过登录态预检")
    sp.set_defaults(func=cmd_collect)

    sp = sub.add_parser("selftest", help="不联网自测解析与导出")
    sp.add_argument("--out", default="selftest_out.csv")
    sp.set_defaults(func=cmd_selftest)

    return p


def main() -> int:
    args = build_parser().parse_args()
    if args.cmd == "selftest":
        return cmd_selftest(args)
    try:
        return asyncio.run(args.func(args))
    except KeyboardInterrupt:
        return 130


if __name__ == "__main__":
    sys.exit(main())
